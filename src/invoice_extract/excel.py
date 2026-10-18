"""Spreadsheet output: one row per invoice, review rows highlighted."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .evaluate import Evaluation
from .schema import FIELDS, MONEY_FIELDS, ExtractionResult

_REVIEW_FILL = PatternFill("solid", fgColor="FFF4CE")
_HEADER_FONT = Font(bold=True)


def write_workbook(
    path: Path, results: list[ExtractionResult], evaluation: Evaluation | None = None
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "invoices"
    header = ["file", *FIELDS, "needs_review", "review_reasons"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    for res in results:
        row: list[object] = [Path(res.source).name]
        for f in FIELDS:
            value = getattr(res.invoice, f)
            row.append(float(value) if f in MONEY_FIELDS and value is not None else value)
        row += ["yes" if res.needs_review else "no", "; ".join(res.review_reasons)]
        ws.append(row)
        if res.needs_review:
            for cell in ws[ws.max_row]:
                cell.fill = _REVIEW_FILL
    for idx, f in enumerate(header, start=1):
        if f in MONEY_FIELDS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = "#,##0.00"
        if f == "invoice_date":
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = "yyyy-mm-dd"
        ws.column_dimensions[get_column_letter(idx)].width = max(12, min(40, len(f) + 6))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    if evaluation is not None:
        acc = wb.create_sheet("accuracy")
        acc.append(["field", "accuracy"])
        for cell in acc[1]:
            cell.font = _HEADER_FONT
        for f, a in evaluation.field_accuracy().items():
            acc.append([f, a])
            acc.cell(acc.max_row, 2).number_format = "0%"
        acc.append([])
        acc.append(["documents fully correct", evaluation.doc_accuracy])
        acc.append(["errors caught by review", evaluation.errors_caught])
        acc.append(["auto-accepted with error", evaluation.auto_accepted_error_rate])
        for row in acc.iter_rows(min_row=acc.max_row - 2, min_col=2, max_col=2):
            row[0].number_format = "0%"
        acc.column_dimensions["A"].width = 28
    wb.save(path)
