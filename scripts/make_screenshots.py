"""Regenerate the README screenshots (needs poppler's pdftoppm and a headless Chrome)."""

from __future__ import annotations

import html
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
CHROME = shutil.which("google-chrome") or shutil.which("chromium") or "chromium"


def invoice_png(pdf: Path, name: str) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        prefix = Path(tmp) / "p"
        subprocess.run(
            ["pdftoppm", "-r", "80", "-png", "-singlefile", str(pdf), str(prefix)], check=True
        )
        img = Image.open(f"{prefix}.png")
        img.crop((0, 0, img.width, int(img.height * 0.42))).save(SHOTS / f"{name}.png")


def chrome_png(page: Path, name: str, size: str) -> None:
    subprocess.run(
        [
            CHROME,
            "--headless=new",
            "--no-sandbox",
            "--disable-gpu",
            "--hide-scrollbars",
            f"--window-size={size}",
            f"--screenshot={SHOTS / name}.png",
            page.as_uri(),
        ],
        check=True,
        capture_output=True,
    )


def spreadsheet_preview(xlsx: Path, out_html: Path, rows: int = 14) -> None:
    ws = load_workbook(xlsx)["invoices"]
    cells = list(ws.iter_rows(values_only=True))
    keep = [0, 1, 2, 3, 6, 8, 9, 10]  # file, number, date, supplier, currency, tax, total, review
    picked = (
        cells[:1]
        + [r for r in cells[1:] if r[10] == "no"][: rows - 4]
        + [r for r in cells[1:] if r[10] == "yes"][:3]
    )
    css = (
        "body{font:13px Calibri,Arial,sans-serif;margin:16px}table{border-collapse:collapse}"
        "td,th{border:1px solid #d0d0d0;padding:4px 8px;white-space:nowrap}"
        "th{background:#f3f3f3}tr.r td{background:#fff4ce}"
    )
    out = [f"<meta charset='utf-8'><style>{css}</style><table>"]
    for i, row in enumerate(picked):
        tag = "th" if i == 0 else "td"
        cls = " class='r'" if i and row[10] == "yes" else ""
        vals = [row[k] for k in keep]
        texts = [
            "" if v is None else str(v)[:10] if k == 2 and i else str(v)
            for k, v in zip(keep, vals, strict=True)
        ]
        cells_html = "".join(f"<{tag}>{html.escape(t)}</{tag}>" for t in texts)
        out.append(f"<tr{cls}>{cells_html}</tr>")
    out_html.write_text("".join(out) + "</table>", encoding="utf-8")


def main() -> None:
    SHOTS.mkdir(parents=True, exist_ok=True)
    for pdf, name in [
        ("classic_es_00", "invoice_classic_es"),
        ("modern_en_00", "invoice_modern_en"),
        ("held_out_vendor_00", "invoice_held_out"),
    ]:
        invoice_png(ROOT / "examples" / f"{pdf}.pdf", name)
    chrome_png(ROOT / "sample_output" / "report.html", "report", "1000,1080")
    with tempfile.TemporaryDirectory() as tmp:
        page = Path(tmp) / "sheet.html"
        spreadsheet_preview(ROOT / "sample_output" / "extracted.xlsx", page)
        chrome_png(page, "spreadsheet", "870,450")


if __name__ == "__main__":
    main()
